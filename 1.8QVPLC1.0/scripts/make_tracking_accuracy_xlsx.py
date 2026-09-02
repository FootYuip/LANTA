# -*- coding: utf-8 -*-
"""Generate tracking accuracy calculation spreadsheet."""

from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
except ImportError:
    import subprocess
    import sys

    subprocess.check_call([sys.executable, "-m", "pip", "install", "openpyxl", "-q"])
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, Border, Side, PatternFill


def load_rows(txt_path: Path) -> list[tuple[float, float]]:
    rows: list[tuple[float, float]] = []
    if not txt_path.exists():
        return rows
    for ln in txt_path.read_text(encoding="utf-8", errors="replace").splitlines():
        parts = ln.strip().replace(",", " ").split()
        if len(parts) < 2:
            continue
        try:
            rows.append((float(parts[0]), float(parts[1])))
        except ValueError:
            continue
    return rows


def main() -> None:
    desktop = Path.home() / "Desktop"
    txt_path = desktop / "跟踪精度.txt"
    out_path = desktop / "跟踪精度计算表.xlsx"
    rows = load_rows(txt_path)

    wb = Workbook()
    ws = wb.active
    ws.title = "跟踪精度计算"

    thin = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(bold=True, color="FFFFFF")
    result_fill = PatternFill("solid", fgColor="FFF2CC")
    formula_fill = PatternFill("solid", fgColor="E2EFDA")
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws.merge_cells("A1:G1")
    ws["A1"] = "天线跟踪精度计算表"
    ws["A1"].font = Font(bold=True, size=14)

    ws.merge_cells("A2:G2")
    ws["A2"] = (
        "公式: δA=√[Σ(Ai-Ā)²/n]   δE=√[Σ(Ei-Ē)²/n]   合成=√(δA²+δE²)   "
        "（第1列方位、第2列俯仰；分母为n）"
    )
    ws["A2"].alignment = Alignment(wrap_text=True)

    headers = ["序号i", "方位 Ai (°)", "俯仰 Ei (°)", "Ai-Ā", "Ei-Ē", "(Ai-Ā)²", "(Ei-Ē)²"]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(4, col, h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center
        cell.border = thin

    max_rows = 50
    data_start = 5
    data_end = data_start + max_rows - 1
    sum_row = 56

    for i in range(max_rows):
        r = data_start + i
        cell_i = ws.cell(r, 1, i + 1)
        cell_i.border = thin
        cell_i.alignment = center
        for c in range(2, 8):
            cell = ws.cell(r, c)
            cell.border = thin
            cell.alignment = center
            cell.number_format = "0.000000"

    for i, (a, e) in enumerate(rows[:max_rows]):
        r = data_start + i
        ws.cell(r, 2, a)
        ws.cell(r, 3, e)

    # Summary block
    ws.cell(sum_row, 1, "样本数 n").font = Font(bold=True)
    ws.cell(sum_row, 2, f"=COUNTA(B{data_start}:B{data_end})")
    ws.cell(sum_row, 2).fill = result_fill
    ws.cell(sum_row, 2).border = thin

    ws.cell(sum_row + 1, 1, "方位均值 Ā").font = Font(bold=True)
    ws.cell(sum_row + 1, 2, f'=IF(B{sum_row}=0,"",AVERAGE(B{data_start}:B{data_end}))')
    ws.cell(sum_row + 1, 2).fill = result_fill
    ws.cell(sum_row + 1, 2).border = thin
    ws.cell(sum_row + 1, 2).number_format = "0.000000"

    ws.cell(sum_row + 2, 1, "俯仰均值 Ē").font = Font(bold=True)
    ws.cell(sum_row + 2, 2, f'=IF(B{sum_row}=0,"",AVERAGE(C{data_start}:C{data_end}))')
    ws.cell(sum_row + 2, 2).fill = result_fill
    ws.cell(sum_row + 2, 2).border = thin
    ws.cell(sum_row + 2, 2).number_format = "0.000000"

    mean_a = f"$B${sum_row + 1}"
    mean_e = f"$B${sum_row + 2}"
    for i in range(max_rows):
        r = data_start + i
        ws.cell(r, 4, f'=IF(B{r}="","",B{r}-{mean_a})')
        ws.cell(r, 5, f'=IF(C{r}="","",C{r}-{mean_e})')
        ws.cell(r, 6, f'=IF(D{r}="","",D{r}^2)')
        ws.cell(r, 7, f'=IF(E{r}="","",E{r}^2)')
        for c in range(4, 8):
            cell = ws.cell(r, c)
            cell.fill = formula_fill
            cell.number_format = "0.0000000000"
            cell.border = thin

    ws.cell(sum_row + 4, 1, "Σ(Ai-Ā)²").font = Font(bold=True)
    ws.cell(sum_row + 4, 2, f"=SUM(F{data_start}:F{data_end})")
    ws.cell(sum_row + 4, 2).fill = formula_fill
    ws.cell(sum_row + 4, 2).border = thin
    ws.cell(sum_row + 4, 2).number_format = "0.0000000000"

    ws.cell(sum_row + 5, 1, "Σ(Ei-Ē)²").font = Font(bold=True)
    ws.cell(sum_row + 5, 2, f"=SUM(G{data_start}:G{data_end})")
    ws.cell(sum_row + 5, 2).fill = formula_fill
    ws.cell(sum_row + 5, 2).border = thin
    ws.cell(sum_row + 5, 2).number_format = "0.0000000000"

    ws.cell(sum_row + 7, 1, "方位精度 δA (°)").font = Font(bold=True, size=12)
    ws.cell(sum_row + 7, 2, f'=IF(B{sum_row}=0,"",SQRT(B{sum_row + 4}/B{sum_row}))')
    ws.cell(sum_row + 7, 2).fill = result_fill
    ws.cell(sum_row + 7, 2).border = thin
    ws.cell(sum_row + 7, 2).font = Font(bold=True, size=12, color="C00000")
    ws.cell(sum_row + 7, 2).number_format = "0.00000"

    ws.cell(sum_row + 8, 1, "俯仰精度 δE (°)").font = Font(bold=True, size=12)
    ws.cell(sum_row + 8, 2, f'=IF(B{sum_row}=0,"",SQRT(B{sum_row + 5}/B{sum_row}))')
    ws.cell(sum_row + 8, 2).fill = result_fill
    ws.cell(sum_row + 8, 2).border = thin
    ws.cell(sum_row + 8, 2).font = Font(bold=True, size=12, color="C00000")
    ws.cell(sum_row + 8, 2).number_format = "0.00000"

    ws.cell(sum_row + 9, 1, "合成精度 √(δA²+δE²) (°)").font = Font(bold=True, size=12)
    ws.cell(
        sum_row + 9,
        2,
        f'=IF(OR(B{sum_row + 7}="",B{sum_row + 8}=""),"",SQRT(B{sum_row + 7}^2+B{sum_row + 8}^2))',
    )
    ws.cell(sum_row + 9, 2).fill = result_fill
    ws.cell(sum_row + 9, 2).border = thin
    ws.cell(sum_row + 9, 2).font = Font(bold=True, size=12, color="C00000")
    ws.cell(sum_row + 9, 2).number_format = "0.00000"

    ws.cell(sum_row + 11, 1, "使用说明").font = Font(bold=True)
    ws.merge_cells(f"A{sum_row + 12}:G{sum_row + 14}")
    ws.cell(sum_row + 12, 1).value = (
        "1. 在 B、C 列填写方位/俯仰角度（可覆盖示例数据，最多50组）；空行不参与计算。\n"
        "2. D~G 列为自动计算，勿改。下方黄色单元格为最终指标。\n"
        "3. 公式与指标要求一致：总体均方根，分母为样本数 n。"
    )
    ws.cell(sum_row + 12, 1).alignment = Alignment(wrap_text=True, vertical="top")

    ws.column_dimensions["A"].width = 28
    for col in "BCDEFG":
        ws.column_dimensions[col].width = 14
    ws.column_dimensions["F"].width = 16
    ws.column_dimensions["G"].width = 16
    ws.row_dimensions[1].height = 22
    ws.row_dimensions[2].height = 36
    ws.row_dimensions[sum_row + 12].height = 60

    wb.save(out_path)
    print(f"saved: {out_path}")
    print(f"prefilled rows: {len(rows)}")


if __name__ == "__main__":
    main()
